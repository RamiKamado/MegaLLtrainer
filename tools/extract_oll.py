# Extract OLL algorithm notes (comments1.xml) paired with the case name below each image cell, plus group headers.
import sys, json, re
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import coordinate_to_tuple
X = sys.argv[1]
ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
root = ET.parse(X + '/xl/comments1.xml').getroot()
wb = openpyxl.load_workbook(sys.argv[2]); ws = wb['OLL']
out = []
for c in root.find('m:commentList', ns):
    ref = c.get('ref'); txt = ''.join(t.text or '' for t in c.iter('{%s}t' % ns['m']))
    r, col = coordinate_to_tuple(ref)
    below = ws.cell(r + 1, col).value
    # group header: walk left on rows 2/3 to the nearest set number / title
    g = None; gt = None
    for cc in range(col, 0, -1):
        if ws.cell(2, cc).value is not None: g = ws.cell(2, cc).value; gt = ws.cell(3, cc).value; break
    out.append({'ref': ref, 'row': r, 'col': col, 'name': below, 'set': g, 'title': gt, 'note': txt})
json.dump(out, open(sys.argv[3], 'w'), ensure_ascii=False, indent=1)
print(len(out), 'notes')
