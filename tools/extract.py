# Extract PLL algorithms from cell notes of the PLL sheet (comments2.xml), pairing image cells with the case name below.
import re, sys, json, html
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import coordinate_to_tuple, get_column_letter
X = sys.argv[1]
ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
root = ET.parse(X + '/xl/comments2.xml').getroot()
notes = {}
for c in root.find('m:commentList', ns):
    ref = c.get('ref')
    txt = ''.join(t.text or '' for t in c.iter('{%s}t' % ns['m']))
    notes[ref] = txt
# threaded comments (discussions)
tns = {'t': 'http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments'}
import os
troot = ET.parse(X + '/xl/threadedComments/threadedComment2.xml').getroot() if os.path.exists(X + '/xl/threadedComments/threadedComment2.xml') else []
threads = {}
for tc in troot:
    ref = tc.get('ref'); txt = ''.join(tc.find('t:text', tns).itertext()) if tc.find('t:text', tns) is not None else ''
    threads.setdefault(ref, []).append(txt)
wb = openpyxl.load_workbook(sys.argv[2])
ws = wb['PLL']
out = []
for ref, txt in notes.items():
    r, c = coordinate_to_tuple(ref)
    below = ws.cell(r + 1, c).value
    val = ws.cell(r, c).value
    out.append({'ref': ref, 'name': below, 'img': bool(val and 'IMAGE' in str(val)), 'note': txt, 'thread': threads.get(ref)})
json.dump(out, open(sys.argv[3], 'w'), ensure_ascii=False, indent=1)
print(len(out), 'notes;', sum(1 for o in out if o['thread']), 'with threads')
